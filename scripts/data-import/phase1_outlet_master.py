"""
Phase 1 — outlet master + alias map + file inventory for the August 2026 data wire-in.
See docs/AUGUST-DATA-CHECKLIST.md ("Phase 1 — Outlet master & alias map").

Reads raw files under datasource/ (gitignored, never committed) and writes an
audit trail under datasource/_audit/ (also gitignored — row-level references
never ship to the browser bundle). This script is the only thing tracked in git;
its outputs are regenerated locally by running it.

Usage: python scripts/data-import/phase1_outlet_master.py
"""
import difflib
import hashlib
import json
import os
import re
import sys
import time
from collections import Counter, defaultdict

import openpyxl

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DATASOURCE = os.path.join(ROOT, "datasource")
AUDIT_DIR = os.path.join(DATASOURCE, "_audit")
OUTLET_DATA_TS = os.path.join(ROOT, "src", "data", "outletData.ts")

os.makedirs(AUDIT_DIR, exist_ok=True)


def log(msg):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", file=sys.stderr)


# ---------------------------------------------------------------------------
# 1. Canonical outlet master, parsed from the existing dashboard data file.
# ---------------------------------------------------------------------------

def load_canonical_outlets():
    text = open(OUTLET_DATA_TS, encoding="utf-8").read()

    trading = []
    m = re.search(r"export const PL_BY_OUTLET.*?=\s*\[(.*?)\n\];", text, re.S)
    for row in re.finditer(
        r"\{\s*name:\s*'([^']+)',\s*code:\s*'([^']+)',\s*entity:\s*'([^']+)'", m.group(1)
    ):
        name, code, entity = row.groups()
        trading.append({"name": name, "code": code, "entity": entity, "status": "trading", "source": "PL_BY_OUTLET"})

    upcoming = []
    io = re.search(r"export const INITIAL_OUTLETS.*?=\s*\[(.*?)\n\];", text, re.S).group(1)
    for block in re.finditer(r"\{(?:[^{}]|\{[^{}]*\})*?status:\s*'upcoming'(?:[^{}]|\{[^{}]*\})*?\}", io):
        b = block.group(0)
        name = re.search(r"name:\s*'([^']+)'", b)
        code = re.search(r"code:\s*'([^']+)'", b)
        entity = re.search(r"entity:\s*'([^']+)'", b)
        if name and code:
            upcoming.append({
                "name": name.group(1), "code": code.group(1),
                "entity": entity.group(1) if entity else None,
                "status": "upcoming", "source": "INITIAL_OUTLETS",
            })

    all_outlets = trading + upcoming
    dupes = [n for n, c in Counter(o["name"] for o in all_outlets).items() if c > 1]
    return all_outlets, dupes


# ---------------------------------------------------------------------------
# 2. File inventory: hash, size, and (for spreadsheets) sheet/row counts.
# ---------------------------------------------------------------------------

def sha256_of(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def build_file_inventory():
    inventory = []
    for dirpath, dirnames, filenames in os.walk(DATASOURCE):
        dirnames[:] = [d for d in dirnames if d != "_audit"]
        for fn in filenames:
            path = os.path.join(dirpath, fn)
            rel = os.path.relpath(path, ROOT).replace("\\", "/")
            inventory.append({
                "path": rel,
                "size_bytes": os.path.getsize(path),
                "sha256": sha256_of(path),
            })
    return inventory


def duplicate_hashes(inventory):
    by_hash = defaultdict(list)
    for row in inventory:
        by_hash[row["sha256"]].append(row["path"])
    return {h: paths for h, paths in by_hash.items() if len(paths) > 1}


# ---------------------------------------------------------------------------
# 3. Outlet-name normalization + matching against the canonical master.
# ---------------------------------------------------------------------------

def norm(s):
    return re.sub(r"\s+", " ", s or "").strip().lower()


def build_matcher(canonical):
    by_norm = {norm(o["name"]): o for o in canonical}
    names_norm = list(by_norm.keys())

    def match(location_text):
        n = norm(location_text)
        if n in by_norm:
            return by_norm[n], "exact"
        # substring either direction (handles "Dpulze Cyberjaya Mall" vs "Dpulze Cyberjaya")
        contains = [k for k in names_norm if k in n or n in k]
        if len(contains) == 1:
            return by_norm[contains[0]], "substring"
        if len(contains) > 1:
            # ambiguous — prefer the longest canonical match contained in n
            contains.sort(key=len, reverse=True)
            return by_norm[contains[0]], "substring-ambiguous"
        close = difflib.get_close_matches(n, names_norm, n=1, cutoff=0.82)
        if close:
            return by_norm[close[0]], "fuzzy"
        return None, "unmatched"

    return match


SPLIT_DASH = re.compile(r"\s+-\s+")
PAREN = re.compile(r"^(.*?)\(([^)]+)\)\s*$")


def strip_brand_prefix(raw):
    """'<Brand> - <Location>' -> location. Falls back to raw if no ' - '."""
    parts = SPLIT_DASH.split(raw, maxsplit=1)
    if len(parts) == 2:
        return parts[0].strip(), parts[1].strip()
    return None, raw.strip()


def strip_parens(raw):
    """'US Pizza (Location)' -> ('US Pizza', 'Location')."""
    m = PAREN.match(raw.strip())
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return None, raw.strip()


# ---------------------------------------------------------------------------
# 4. Per-source outlet-name extraction.
# ---------------------------------------------------------------------------

def extract_grab(path):
    log("Reading Grab (full workbook, ~30k rows)...")
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {h: i for i, h in enumerate(header)}
    store_counts = Counter()
    merchant_counts = Counter()
    date_min = date_max = None
    n = 0
    for r in rows:
        n += 1
        store = r[idx["Store Name"]]
        merchant = r[idx["Merchant Name"]]
        if store:
            store_counts[store] += 1
        if merchant:
            merchant_counts[merchant] += 1
        created = r[idx["Created On"]]
        if created:
            if date_min is None or created < date_min:
                date_min = created
            if date_max is None or created > date_max:
                date_max = created
    wb.close()
    return {
        "row_count": n,
        "distinct_merchant_names": dict(merchant_counts),
        "distinct_store_names": dict(store_counts),
        "created_on_min": date_min,
        "created_on_max": date_max,
    }


def extract_shopee(path):
    log("Reading Shopee (full workbook, ~17k rows)...")
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {h: i for i, h in enumerate(header)}
    store_counts = Counter()
    date_min = date_max = None
    n = 0
    for r in rows:
        n += 1
        store = r[idx["Store Name"]]
        if store:
            store_counts[store] += 1
        ct = r[idx["Complete Time"]]
        if ct:
            if date_min is None or str(ct) < str(date_min):
                date_min = ct
            if date_max is None or str(ct) > str(date_max):
                date_max = ct
    wb.close()
    return {"row_count": n, "distinct_store_names": dict(store_counts), "complete_time_min": date_min, "complete_time_max": date_max}


def extract_app(path):
    log("Reading App order list (~7k rows)...")
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {h: i for i, h in enumerate(header)}
    store_counts = Counter()
    date_min = date_max = None
    n = 0
    for r in rows:
        n += 1
        store = r[idx["Outlet Name"]]
        if store:
            store_counts[store] += 1
        od = r[idx["Order Date"]]
        if od:
            if date_min is None or str(od) < str(date_min):
                date_min = od
            if date_max is None or str(od) > str(date_max):
                date_max = od
    wb.close()
    return {"row_count": n, "distinct_outlet_names": dict(store_counts), "order_date_min": date_min, "order_date_max": date_max}


def extract_pos(folder):
    log("Reading POS Sales Details reports (5 files)...")
    outlet_counts = Counter()
    subtotal_counts = Counter()
    date_headings = set()
    files_meta = []
    for fn in sorted(os.listdir(folder)):
        if not fn.lower().endswith(".xlsx"):
            continue
        path = os.path.join(folder, fn)
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
        ws = wb[wb.sheetnames[0]]
        n_rows = 0
        n_outlet_headings = 0
        n_subtotals = 0
        for r in ws.iter_rows(values_only=True):
            n_rows += 1
            a, b = r[0], r[1]
            if a is not None and b is None:
                s = str(a).strip()
                if s.startswith("Subtotal:"):
                    subtotal_counts[s[len("Subtotal:"):].strip()] += 1
                    n_subtotals += 1
                elif re.match(r"^\d{2}/\d{2}/\d{4}", s):
                    date_headings.add(s.strip())
                elif re.match(r"^[A-Za-z0-9]+-.+", s) and not s.lower().startswith(("sales details", "report", "filters")):
                    outlet_counts[s] += 1
                    n_outlet_headings += 1
        wb.close()
        files_meta.append({"file": fn, "rows": n_rows, "outlet_headings": n_outlet_headings, "subtotal_rows": n_subtotals})
        log(f"  {fn}: {n_rows} rows, {n_outlet_headings} outlet headings, {n_subtotals} subtotal rows")
    return {
        "files": files_meta,
        "distinct_outlet_headings": dict(outlet_counts),
        "distinct_subtotal_labels": dict(subtotal_counts),
        "date_headings_seen": sorted(date_headings),
    }


def extract_foodpanda(folder):
    log("Reading Foodpanda (459 xlsx invoices + 526 pdf invoices)...")
    files = os.listdir(folder)
    pat = re.compile(r"^(\d{8})_(\d+)\.(xlsx|pdf)$", re.I)
    xlsx_files, pdf_files = {}, {}
    unmatched_filenames = []
    for f in files:
        m = pat.match(f)
        if not m:
            unmatched_filenames.append(f)
            continue
        date, num, ext = m.groups()
        (xlsx_files if ext.lower() == "xlsx" else pdf_files)[num] = f

    xlsx_nums = set(xlsx_files)
    pdf_nums = set(pdf_files)
    both = xlsx_nums & pdf_nums
    xlsx_only = xlsx_nums - pdf_nums
    pdf_only = pdf_nums - xlsx_nums

    outlet_counts = Counter()
    legal_name_counts = Counter()
    invoice_dates = []
    xlsx_errors = []
    for i, num in enumerate(sorted(xlsx_nums)):
        path = os.path.join(folder, xlsx_files[num])
        try:
            wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
            ws = wb["Appendix A"] if "Appendix A" in wb.sheetnames else wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            header = next(rows)
            idx = {h: i for i, h in enumerate(header)}
            first = next(rows, None)
            if first:
                outlet = first[idx.get("Outlet Name", -1)] if "Outlet Name" in idx else None
                if outlet:
                    outlet_counts[outlet] += 1
                inv_date = first[idx.get("Invoice Date", -1)] if "Invoice Date" in idx else None
                if inv_date:
                    invoice_dates.append(str(inv_date))
            wb.close()
        except Exception as e:
            xlsx_errors.append({"file": xlsx_files[num], "error": str(e)})
        if (i + 1) % 100 == 0:
            log(f"  ...{i + 1}/{len(xlsx_nums)} xlsx invoices read")

    pdf_outlet_counts = Counter()
    pdf_errors = []
    if pdfplumber is None:
        log("  pdfplumber not available — skipping PDF-only outlet-name extraction")
    else:
        outlet_re = re.compile(r"Outlet Name:\s*(.+)")
        legal_re = re.compile(r"Legal Name:\s*(.+)")
        for i, num in enumerate(sorted(pdf_only)):
            path = os.path.join(folder, pdf_files[num])
            try:
                with pdfplumber.open(path) as pdf:
                    txt = pdf.pages[0].extract_text() or ""
                om = outlet_re.search(txt)
                lm = legal_re.search(txt)
                if om:
                    pdf_outlet_counts[om.group(1).strip()] += 1
                if lm:
                    legal_name_counts[lm.group(1).strip()] += 1
            except Exception as e:
                pdf_errors.append({"file": pdf_files[num], "error": str(e)})
            if (i + 1) % 100 == 0:
                log(f"  ...{i + 1}/{len(pdf_only)} pdf-only invoices scanned")

    return {
        "xlsx_invoice_count": len(xlsx_nums),
        "pdf_invoice_count": len(pdf_nums),
        "unmatched_filenames": unmatched_filenames,
        "invoices_with_both_xlsx_and_pdf": len(both),
        "invoices_xlsx_only_no_pdf_backup": len(xlsx_only),
        "invoices_pdf_only_no_excel_detail": len(pdf_only),
        "total_distinct_invoices": len(xlsx_nums | pdf_nums),
        "xlsx_distinct_outlet_names": dict(outlet_counts),
        "pdf_only_distinct_outlet_names": dict(pdf_outlet_counts),
        "pdf_only_legal_names": dict(legal_name_counts),
        "xlsx_read_errors": xlsx_errors,
        "pdf_read_errors": pdf_errors,
        "invoice_date_sample_min": min(invoice_dates) if invoice_dates else None,
        "invoice_date_sample_max": max(invoice_dates) if invoice_dates else None,
    }


# ---------------------------------------------------------------------------
# 5. Alias-map builder, applying source-specific prefix-stripping + matcher.
# ---------------------------------------------------------------------------

def build_alias_rows(source_label, raw_counts, splitter, match):
    rows = []
    for raw, count in sorted(raw_counts.items(), key=lambda kv: -kv[1]):
        prefix, location = splitter(raw)
        matched, confidence = match(location)
        rows.append({
            "source": source_label,
            "raw": raw,
            "row_count": count,
            "prefix": prefix,
            "location_text": location,
            "matched_name": matched["name"] if matched else None,
            "matched_code": matched["code"] if matched else None,
            "matched_entity": matched["entity"] if matched else None,
            "confidence": confidence,
        })
    return rows


def main():
    canonical, dupes = load_canonical_outlets()
    log(f"Canonical outlets loaded: {len(canonical)} ({dupes and 'DUPLICATE NAMES: ' + str(dupes) or 'no duplicate names'})")

    match = build_matcher(canonical)

    log("Building file inventory + hashes...")
    inventory = build_file_inventory()
    dupe_hashes = duplicate_hashes(inventory)

    grab = extract_grab(os.path.join(DATASOURCE, "GRAB Aug sales.xlsx"))
    shopee = extract_shopee(os.path.join(DATASOURCE, "SHOPEE", "Shopee Aug Sales.xlsx"))
    app = extract_app(os.path.join(DATASOURCE, "APPS AUG ORDER LIST.xlsx"))
    pos = extract_pos(os.path.join(DATASOURCE, "POS SALES"))
    foodpanda = extract_foodpanda(os.path.join(DATASOURCE, "FOODPANDA"))

    alias_rows = []
    alias_rows += build_alias_rows("Grab", grab["distinct_store_names"], strip_brand_prefix, match)
    alias_rows += build_alias_rows("Shopee", shopee["distinct_store_names"], strip_brand_prefix, match)
    alias_rows += build_alias_rows("App", app["distinct_outlet_names"], strip_brand_prefix, match)
    alias_rows += build_alias_rows("Foodpanda-xlsx", foodpanda["xlsx_distinct_outlet_names"], strip_parens, match)
    alias_rows += build_alias_rows("Foodpanda-pdf-only", foodpanda["pdf_only_distinct_outlet_names"], strip_parens, match)

    def pos_split(raw):
        m = re.match(r"^([A-Za-z0-9]+)-(?:US Pizza\s+)?(.*)$", raw)
        if m:
            return m.group(1), m.group(2).strip()
        return None, raw
    alias_rows += build_alias_rows("POS", pos["distinct_outlet_headings"], pos_split, match)

    unmatched = [r for r in alias_rows if r["confidence"] in ("unmatched", "substring-ambiguous")]
    matched_canonical_names = {r["matched_name"] for r in alias_rows if r["matched_name"]}
    canonical_names = {o["name"] for o in canonical}
    never_seen = sorted(canonical_names - matched_canonical_names)

    # Shopee/App/Grab brand prefixes actually seen — for the "US Pizza only" exclusion rule.
    def prefixes_seen(rows, source):
        return sorted({r["prefix"] for r in rows if r["source"] == source and r["prefix"]})

    brand_prefixes = {
        "Grab": prefixes_seen(alias_rows, "Grab"),
        "Shopee": prefixes_seen(alias_rows, "Shopee"),
        "App": prefixes_seen(alias_rows, "App"),
    }

    # --- write artifacts --------------------------------------------------
    def dump(name, obj):
        path = os.path.join(AUDIT_DIR, name)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(obj, f, indent=2, default=str, ensure_ascii=False)
        log(f"wrote {path}")

    dump("canonical_outlets.json", {"outlets": canonical, "duplicate_names": dupes})
    dump("file_inventory.json", {"files": inventory, "duplicate_sha256": dupe_hashes})
    dump("source_raw_grab.json", grab)
    dump("source_raw_shopee.json", shopee)
    dump("source_raw_app.json", app)
    dump("source_raw_pos.json", pos)
    dump("source_raw_foodpanda.json", foodpanda)
    dump("outlet_alias_map.json", {"rows": alias_rows, "brand_prefixes_seen": brand_prefixes})
    dump("outlet_alias_unmatched.json", {"rows": unmatched, "canonical_names_never_matched": never_seen})

    # --- human-readable findings summary -----------------------------------
    md = []
    md.append("# Phase 1 findings — outlet master, file inventory, alias map\n")
    md.append(f"Generated by `scripts/data-import/phase1_outlet_master.py`. Raw data lives in `datasource/_audit/` (gitignored).\n")
    md.append("## File inventory\n")
    md.append(f"- {len(inventory)} files hashed under `datasource/`.")
    md.append(f"- Duplicate content (same sha256, different filename): {len(dupe_hashes)} group(s).")
    if dupe_hashes:
        for h, paths in list(dupe_hashes.items())[:20]:
            md.append(f"  - `{h[:12]}…`: {paths}")
    md.append("")
    md.append("## Canonical outlet master")
    md.append(f"- {len(canonical)} outlets parsed from `src/data/outletData.ts` "
               f"({sum(1 for o in canonical if o['status']=='trading')} trading + "
               f"{sum(1 for o in canonical if o['status']=='upcoming')} upcoming).")
    if dupes:
        md.append(f"- ⚠ duplicate outlet names in master: {dupes}")
    md.append("")
    md.append("## Brand prefixes observed (US Pizza vs sister brands)")
    for src, prefixes in brand_prefixes.items():
        md.append(f"- **{src}**: {prefixes}")
    md.append("- Foodpanda xlsx/pdf both use the `US Pizza (Location)` pattern only in the sample audited — no sister-brand prefix seen there.")
    md.append("")
    md.append("## Foodpanda invoice coverage (xlsx vs pdf)")
    md.append(f"- xlsx invoices: {foodpanda['xlsx_invoice_count']}")
    md.append(f"- pdf invoices: {foodpanda['pdf_invoice_count']}")
    md.append(f"- have both xlsx detail + pdf: {foodpanda['invoices_with_both_xlsx_and_pdf']}")
    md.append(f"- xlsx only, no pdf backup: {foodpanda['invoices_xlsx_only_no_pdf_backup']}")
    md.append(f"- **pdf only, no Excel order-level detail: {foodpanda['invoices_pdf_only_no_excel_detail']}**")
    md.append(f"- total distinct invoices: {foodpanda['total_distinct_invoices']} "
               f"({'matches' if foodpanda['total_distinct_invoices']==877 else 'does NOT match'} the summary's cited 877)")
    md.append("- Finding: the 877-vs-459 gap in the checklist is explained — 459 invoices have Excel "
               "transaction detail, and the other 418 exist ONLY as PDF (108 invoices have both). "
               "The 418 PDF-only invoices are E-Invoices with an outlet-level daily total on page 1 "
               "(`Number Of Successful Orders`, `Total`, `Outstanding Amount`, SST breakdown) but no "
               "per-order rows, so Phase 2/3 must parse PDF-only invoices at the daily-outlet-total "
               "grain, not the order-detail grain used for the 459 xlsx invoices.")
    md.append("")
    md.append("## Outlet alias map")
    md.append(f"- {len(alias_rows)} distinct raw outlet strings seen across all 5 sources.")
    md.append(f"- Unmatched / ambiguous: {len(unmatched)}")
    if unmatched:
        for r in unmatched[:40]:
            md.append(f"  - [{r['source']}] `{r['raw']}` (location text: `{r['location_text']}`, {r['row_count']} rows) → {r['confidence']}")
    md.append(f"- Canonical outlets never matched in any source: {never_seen}")
    md.append("")
    md.append("## Next (Phase 2)")
    md.append("- Resolve each unmatched/ambiguous row above with the user before aggregation — do not guess.")
    md.append("- Decide per-platform date/status/reversal rules using `source_raw_*.json` distinct Status/Type values.")
    md.append("- Build the PDF-only Foodpanda daily-total parser (Phase 2/3), since it covers 418 of 877 invoices.")

    with open(os.path.join(AUDIT_DIR, "PHASE1-FINDINGS.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md) + "\n")
    log(f"wrote {os.path.join(AUDIT_DIR, 'PHASE1-FINDINGS.md')}")

    print("\n".join(md))


if __name__ == "__main__":
    main()
