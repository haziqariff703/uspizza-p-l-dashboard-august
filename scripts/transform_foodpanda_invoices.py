#!/usr/bin/env python3
"""
Foodpanda PDF to Excel Invoice Transformer

Transforms Foodpanda / Delivery Hero PDF tax invoices (e-invoices and self-billed
e-invoices) into the standardized `Foodpanda_Invoice_Report.xlsx` (`Invoice Master` sheet)
format as specified in `docs/FOODPANDA_EXCEL_PARSING_SPECIFICATION.md` and required by
`src/lib/salesImportParser.ts`.

Usage:
    python scripts/transform_foodpanda_invoices.py --input-dir "datasource/FOODPANDA/foodpanda_invoice_pdf" --output "Foodpanda_Invoice_Report.xlsx"
"""

import argparse
import concurrent.futures
import datetime
import glob
import os
import re
import sys

try:
    import pdfplumber
except ImportError:
    sys.exit("Error: pdfplumber is required. Run 'pip install pdfplumber'.")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Error: openpyxl is required. Run 'pip install openpyxl'.")


def clean_num(val) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(",", "").replace("RM", "").strip()
    # Normalize special minus/dash symbols
    s = s.replace("–", "-").replace("—", "-")
    if not s or s == "-":
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def parse_date(date_str: str) -> str:
    if not date_str:
        return ""
    m = re.search(r"(\d{4})[./-](\d{2})[./-](\d{2})", date_str)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m2 = re.search(r"(\d{2})[./-](\d{2})[./-](\d{4})", date_str)
    if m2:
        return f"{m2.group(3)}-{m2.group(2)}-{m2.group(1)}"
    return date_str.strip()


def parse_single_pdf(pdf_path: str) -> dict:
    """Extract all relevant financial and audit fields from a Foodpanda invoice PDF."""
    filename = os.path.basename(pdf_path)
    data = {
        "source_file": filename,
        "data_source": "PDF only",
        "business_date": "",
        "invoice_issue_date": "",
        "period_from": "",
        "period_to": "",
        "outlet_name": "",
        "legal_entity": "",
        "invoice_no": "",
        "self_billed_inv_no": "",
        "vendor_id": "",
        "vendor_code": "",
        "vendor_tin": "",
        # Order counts
        "normal_orders": 0,
        "cancelled_partial_orders": 0,
        "cancelled_full_orders": 0,
        "pandabox_orders": 0,
        "dine_in_orders": 0,
        "targeted_customer_orders": 0,
        "total_orders": 0,
        # Revenue
        "normal_revenue": 0.0,
        "cancelled_partial_revenue": 0.0,
        "cancelled_full_revenue": 0.0,
        "pandabox_revenue": 0.0,
        "dine_in_revenue": 0.0,
        "targeted_customer_revenue": 0.0,
        "total_revenue": 0.0,
        "already_received": 0.0,
        "outstanding_a": 0.0,
        # Commission rate
        "commission_rate": 20.0,
        # Charges (8% SST)
        "commission_base": 0.0,
        "commission_sst": 0.0,
        "wastage_commission_base": 0.0,
        "wastage_commission_sst": 0.0,
        "fees_adj_sst_base": 0.0,
        "fees_adj_sst_sst": 0.0,
        "fees_adj_non_sst_base": 0.0,
        "fees_adj_non_sst_sst": 0.0,
        "pandabox_fee_base": 0.0,
        "pandabox_fee_sst": 0.0,
        "targeted_cust_fee_base": 0.0,
        "targeted_cust_fee_sst": 0.0,
        "waiting_time_fee_base": 0.0,
        "waiting_time_fee_sst": 0.0,
        # Totals Page 1
        "total_excl_tax": 0.0,
        "total_sst": 0.0,
        "total_incl_tax_c": 0.0,
        # Self-Billed Page (Sales)
        "sales_value_excl_sst": 0.0,
        "sales_sst_6pct": 0.0,
        "sales_incl_sst": 0.0,
        # Settlement
        "todays_earnings": 0.0,
        "outstanding_prev_f": 0.0,
        "total_payable": 0.0,
        # Appendix counts
        "appendix_c_open_items": 0,
        "error": None,
    }

    try:
        with pdfplumber.open(pdf_path) as pdf:
            pages = pdf.pages
            num_pages = len(pages)
            if num_pages == 0:
                data["error"] = "Empty PDF"
                return data

            p1_text = pages[0].extract_text() or ""
            last_text = pages[-1].extract_text() or ""

            # -------------------------------------------------------------
            # 1. Header Information (Page 1)
            # -------------------------------------------------------------
            m_outlet = re.search(r"Outlet Name:\s*(.+)", p1_text)
            if m_outlet:
                data["outlet_name"] = m_outlet.group(1).strip()

            m_legal = re.search(r"Legal Name:\s*(.+)", p1_text)
            if m_legal:
                data["legal_entity"] = m_legal.group(1).strip()

            m_inv = re.search(r"Invoice Number:\s*(\d+)", p1_text)
            if m_inv:
                data["invoice_no"] = m_inv.group(1).strip()

            m_date = re.search(r"Invoice Date and Time:\s*(\d{4}[./-]\d{2}[./-]\d{2})", p1_text)
            if m_date:
                data["invoice_issue_date"] = parse_date(m_date.group(1))

            m_period = re.search(r"Invoice Period:\s*(\d{4}[./-]\d{2}[./-]\d{2})\s+to\s+(\d{4}[./-]\d{2}[./-]\d{2})", p1_text)
            if m_period:
                data["period_from"] = parse_date(m_period.group(1))
                data["period_to"] = parse_date(m_period.group(2))
                data["business_date"] = data["period_from"]
            else:
                # Fallback to date in filename (YYYYMMDD_...)
                m_fn = re.search(r"^(\d{4})(\d{2})(\d{2})_", filename)
                if m_fn:
                    data["business_date"] = f"{m_fn.group(1)}-{m_fn.group(2)}-{m_fn.group(3)}"

            m_vid = re.search(r"(?:Vendor|Supplier) Identification Number:\s*([A-Za-z0-9\-]+)", p1_text)
            if m_vid:
                data["vendor_id"] = m_vid.group(1).strip()

            m_tin = re.search(r"(?:Vendor|Supplier) TIN:\s*([A-Za-z0-9\-]+)", p1_text)
            if m_tin:
                data["vendor_tin"] = m_tin.group(1).strip()

            # -------------------------------------------------------------
            # 2. Order Counts & Revenue (Page 1)
            # -------------------------------------------------------------
            p1_lines = p1_text.splitlines()
            for line in p1_lines:
                l_str = line.strip()
                # Normal Orders 17 956.29
                m = re.match(r"^Normal Orders\s+(\d+)\s+([\d,.\-]+)", l_str, re.I)
                if m:
                    data["normal_orders"] = int(m.group(1))
                    data["normal_revenue"] = clean_num(m.group(2))
                    continue

                # Cancelled Orders (Partial food reimbursement) 0 0.00
                m = re.match(r"^Cancelled Orders \(Partial.*?\)\s+(\d+)\s+([\d,.\-]+)", l_str, re.I)
                if m:
                    data["cancelled_partial_orders"] = int(m.group(1))
                    data["cancelled_partial_revenue"] = clean_num(m.group(2))
                    continue

                # Cancelled Orders (Full food reimbursement) 1 54.51
                m = re.match(r"^Cancelled Orders \(Full.*?\)\s+(\d+)\s+([\d,.\-]+)", l_str, re.I)
                if m:
                    data["cancelled_full_orders"] = int(m.group(1))
                    data["cancelled_full_revenue"] = clean_num(m.group(2))
                    continue

                # Pandabox Orders 0 0.00
                m = re.match(r"^Pandabox Orders\s+(\d+)\s+([\d,.\-]+)", l_str, re.I)
                if m:
                    data["pandabox_orders"] = int(m.group(1))
                    data["pandabox_revenue"] = clean_num(m.group(2))
                    continue

                # Dine-in Orders 0 0.00
                m = re.match(r"^Dine-in Orders\s+(\d+)\s+([\d,.\-]+)", l_str, re.I)
                if m:
                    data["dine_in_orders"] = int(m.group(1))
                    data["dine_in_revenue"] = clean_num(m.group(2))
                    continue

                # Targeted Customer Orders 0 0.00
                m = re.match(r"^Targeted Customer Orders\s+(\d+)\s+([\d,.\-]+)", l_str, re.I)
                if m:
                    data["targeted_customer_orders"] = int(m.group(1))
                    data["targeted_customer_revenue"] = clean_num(m.group(2))
                    continue

                # Total 18 1,010.80 (in order section)
                m = re.match(r"^Total\s+(\d+)\s+([\d,.\-]+)", l_str)
                if m and data["total_orders"] == 0:
                    data["total_orders"] = int(m.group(1))
                    data["total_revenue"] = clean_num(m.group(2))
                    continue

                # Already Received Amount 0 0.00
                m = re.search(r"Already Received Amount\s+([\d,.\-]+)", l_str)
                if m:
                    data["already_received"] = clean_num(m.group(1))
                    continue

                # Outstanding Amount [a] 0 1,010.80
                m = re.search(r"Outstanding Amount \[a\]\s+(?:\d+\s+)?([\d,.\-]+)", l_str)
                if m:
                    data["outstanding_a"] = clean_num(m.group(1))
                    continue

            # Fallback if total row wasn't picked up
            if data["total_orders"] == 0:
                data["total_orders"] = (
                    data["normal_orders"] + data["cancelled_partial_orders"] +
                    data["cancelled_full_orders"] + data["pandabox_orders"] +
                    data["dine_in_orders"] + data["targeted_customer_orders"]
                )
            if data["total_revenue"] == 0.0:
                data["total_revenue"] = round(
                    data["normal_revenue"] + data["cancelled_partial_revenue"] +
                    data["cancelled_full_revenue"] + data["pandabox_revenue"] +
                    data["dine_in_revenue"] + data["targeted_customer_revenue"], 2
                )
            if data["outstanding_a"] == 0.0 and data["total_revenue"] > 0:
                data["outstanding_a"] = data["total_revenue"] - data["already_received"]

            # -------------------------------------------------------------
            # 3. Platform Charges (Page 1)
            # -------------------------------------------------------------
            in_charges = False
            for line in p1_lines:
                l_str = line.strip()
                if "SST Rate SST Base" in l_str or "Description SST Rate" in l_str:
                    in_charges = True
                    continue
                if in_charges:
                    if "Subtotal [b]" in l_str or "Total (excluding tax)" in l_str:
                        in_charges = False

                    # Match: Description ... SST% Base SST
                    # E.g. Commission 8% 180.42 14.44
                    m = re.match(r"^(.+?)\s+(\d+%)\s+([\d,.\-]+)\s+([\d,.\-]+)$", l_str)
                    if m:
                        desc = m.group(1).strip()
                        base = clean_num(m.group(3))
                        sst = clean_num(m.group(4))

                        if desc == "Commission":
                            data["commission_base"] = base
                            data["commission_sst"] = sst
                        elif "Cancelled Orders" in desc or "Wastage" in desc:
                            data["wastage_commission_base"] = base
                            data["wastage_commission_sst"] = sst
                        elif "Targeted Customer Orders" in desc:
                            data["targeted_cust_fee_base"] = base
                            data["targeted_cust_fee_sst"] = sst
                        elif "Pandabox Fee" in desc:
                            data["pandabox_fee_base"] = base
                            data["pandabox_fee_sst"] = sst
                        elif "Waiting Time" in desc:
                            data["waiting_time_fee_base"] = base
                            data["waiting_time_fee_sst"] = sst
                        elif "Appendix B" in desc or "Fees & Adjustments" in desc:
                            data["fees_adj_sst_base"] = base
                            data["fees_adj_sst_sst"] = sst

                # -------------------------------------------------------------
                # 4. Totals and Settlement on Page 1
                # -------------------------------------------------------------
                m_excl = re.match(r"^Total \(excluding tax\)\s+([\d,.\-]+)", l_str)
                if m_excl:
                    data["total_excl_tax"] = clean_num(m_excl.group(1))

                m_sst = re.match(r"^Total SST On Service\s+([\d,.\-]+)", l_str)
                if m_sst:
                    data["total_sst"] = clean_num(m_sst.group(1))

                m_incl = re.match(r"^Total \(including tax\) \[c\]\s+([\d,.\-]+)", l_str)
                if m_incl:
                    data["total_incl_tax_c"] = clean_num(m_incl.group(1))

                # Today’s Earnings [a-c]
                m_earn = re.search(r"Today[^\w\s]*s Earnings \[a[\-_]c\]\s+([\d,.\-]+)", l_str, re.I)
                if m_earn:
                    data["todays_earnings"] = clean_num(m_earn.group(1))

                # Outstanding Amount From Previous Periods [f] (Appendix C) 466.89
                m_prev = re.search(r"Outstanding Amount From Previous Periods \[f\].*?\s+([\d,.\-]+)", l_str)
                if m_prev:
                    data["outstanding_prev_f"] = clean_num(m_prev.group(1))

                # Outstanding amount [f] + Today’s earnings [a-c] 1,271.72
                m_tot_pay = re.search(r"Outstanding amount \[f\] \+ Today[^\w\s]*s earnings \[a[\-_]c\]\s+([\d,.\-]+)", l_str, re.I)
                if m_tot_pay:
                    data["total_payable"] = clean_num(m_tot_pay.group(1))

            # -------------------------------------------------------------
            # 5. Commission Rate (Page 2 if available)
            # -------------------------------------------------------------
            if num_pages > 1:
                p2_text = pages[1].extract_text() or ""
                m_rate = re.search(r"Commission Rate\s*\n.*?FOODPANDA\s+[\d.]+\s+[\d.]+\s+([\d,.]+)", p2_text, re.S)
                if not m_rate:
                    m_rate = re.search(r"FOODPANDA\s+[\d.]+\s+[\d.]+\s+([\d,.]+)", p2_text)
                if m_rate:
                    data["commission_rate"] = clean_num(m_rate.group(1))
                elif data["commission_base"] > 0 and data["total_revenue"] > 0:
                    calc_rate = round((data["commission_base"] / data["total_revenue"]) * 100, 1)
                    if 15 <= calc_rate <= 35:
                        data["commission_rate"] = calc_rate

            # -------------------------------------------------------------
            # 6. Appendix C : Open Items count
            # -------------------------------------------------------------
            appx_c_count = 0
            for idx in range(1, num_pages - 1):
                page_txt = pages[idx].extract_text() or ""
                if "Appendix C : Open Items" in page_txt:
                    for line in page_txt.splitlines():
                        if re.match(r"^\d{10,14}\s+\d{2}\.\d{2}\.\d{4}", line.strip()):
                            appx_c_count += 1
            data["appendix_c_open_items"] = appx_c_count

            # -------------------------------------------------------------
            # 7. Self-Billed E-Invoice Page (Last Page)
            # -------------------------------------------------------------
            if "SELF-BILL" in last_text.upper():
                m_sb_inv = re.search(r"Invoice Number:\s*([A-Za-z0-9]+)", last_text)
                if m_sb_inv:
                    data["self_billed_inv_no"] = m_sb_inv.group(1).strip()

                last_lines = last_text.splitlines()
                for l_str in last_lines:
                    l_str = l_str.strip()
                    m_ex = re.match(r"^Total \(excluding tax\)\s+([\d,.\-]+)", l_str)
                    if m_ex:
                        data["sales_value_excl_sst"] = clean_num(m_ex.group(1))
                    m_sst6 = re.match(r"^Total SST\s+([\d,.\-]+)", l_str)
                    if m_sst6:
                        data["sales_sst_6pct"] = clean_num(m_sst6.group(1))
                    m_in6 = re.match(r"^Total \(including tax\)\s+([\d,.\-]+)", l_str)
                    if m_in6:
                        data["sales_incl_sst"] = clean_num(m_in6.group(1))

            if data["sales_incl_sst"] == 0.0 and data["total_revenue"] > 0:
                data["sales_incl_sst"] = data["total_revenue"]
            if data["sales_value_excl_sst"] == 0.0 and data["sales_incl_sst"] > 0:
                data["sales_value_excl_sst"] = round(data["sales_incl_sst"] / 1.06, 2)
                data["sales_sst_6pct"] = round(data["sales_incl_sst"] - data["sales_value_excl_sst"], 2)

            if not data["self_billed_inv_no"] and data["invoice_no"]:
                data["self_billed_inv_no"] = f"{data['invoice_no']}S"

    except Exception as e:
        data["error"] = str(e)

    return data


# ===========================================================================
# 8. Excel Workbook Writer
# ===========================================================================

HEADERS_ROW_2 = [
    "Data Source", "Business Date", "Invoice Issue Date", "Period From", "Period To",
    "Outlet Name", "Legal Entity", "Invoice No.", "Self-Billed Inv. No.", "Vendor ID",
    "Vendor Code", "Vendor TIN", "Normal Orders", "Cancelled (Partial Reimb.) Orders",
    "Cancelled (Full Reimb.) Orders", "Pandabox Orders", "Dine-in Orders", "Targeted Customer Orders",
    "Total Orders", "Order Lines (Appx A)", "Normal (RM)", "Cancelled (Partial Reimb.) (RM)",
    "Cancelled (Full Reimb.) (RM)", "Pandabox (RM)", "Dine-in (RM)", "Targeted Customer (RM)",
    "Total Revenue (RM)", "Already Received (RM)", "Outstanding [a] (RM)", "Products Value (RM)",
    "Packaging Fees (RM)", "MOV (RM)", "Vendor Delivery Fee (RM)", "Voucher by Vendor (RM)",
    "Discount by Vendor (RM)", "Pandabox Voucher (RM)", "Delivery Fee Discount (RM)",
    "Customer Targeting Fee (RM)", "Restaurant Revenue (RM)", "Commission Rate (%)",
    "Commission Base (RM)", "Commission SST (RM)", "Wastage Commission Base (RM)",
    "Wastage Commission SST (RM)", "Fees & Adj (SST) Base (RM)", "Fees & Adj (SST) SST (RM)",
    "Fees & Adj (Non-SST) Base (RM)", "Fees & Adj (Non-SST) SST (RM)", "Pandabox Fee Base (RM)",
    "Pandabox Fee SST (RM)", "Targeted Cust. Fee Base (RM)", "Targeted Cust. Fee SST (RM)",
    "Waiting Time Fee Base (RM)", "Waiting Time Fee SST (RM)", "Total Excl. Tax (RM)",
    "Total SST (RM)", "Total Incl. Tax [c] (RM)", "Sales Value Excl. SST (RM)",
    "Sales SST 6% (RM)", "Sales Incl. SST (RM)", "Today's Earnings [a-c] (RM)",
    "Outstanding Prev. [f] (RM)", "Total Payable [f]+[a-c] (RM)", "Appendix B Lines",
    "Appendix B Total (RM)", "Appendix C Open Items", "Source File"
]

GROUP_RANGES_ROW_1 = [
    ("A1:L1", "INVOICE DETAILS", "2F4F4F"),
    ("M1:T1", "ORDER COUNT", "1E3A8A"),
    ("U1:AC1", "REVENUE VIA FOODPANDA", "065F46"),
    ("AD1:AM1", "ORDER-LEVEL DETAIL (PORTAL APPENDIX A)", "4B5563"),
    ("AN1:BE1", "FOODPANDA CHARGES (COMMISSION, FEES & SST)", "991B1B"),
    ("BF1:BH1", "SELF-BILLED E-INVOICE (SALES)", "166534"),
    ("BI1:BK1", "SETTLEMENT", "854D0E"),
    ("BL1:BO1", "REFERENCE", "374151"),
]


def write_excel_report(records: list, output_path: str):
    """Generate the standardized Foodpanda_Invoice_Report.xlsx with Invoice Master."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice Master"
    ws.views.sheetView[0].showGridLines = True

    # Styles
    font_group = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=9, bold=True, color="1E293B")
    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_data = Font(name="Calibri", size=9)
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # 1. Write Row 1 Group Headers
    for cell_range, title, bg_color in GROUP_RANGES_ROW_1:
        ws.merge_cells(cell_range)
        top_left = ws[cell_range.split(":")[0]]
        top_left.value = title
        top_left.font = font_group
        top_left.alignment = align_center
        fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(cell_range)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(r, c).fill = fill

    # 2. Write Row 2 Column Headers
    ws.row_dimensions[2].height = 24
    for col_idx, header in enumerate(HEADERS_ROW_2, start=1):
        cell = ws.cell(2, col_idx, header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    # 3. Write Data Rows
    sorted_records = sorted(records, key=lambda x: (x.get("business_date") or "", x.get("outlet_name") or ""))

    for row_idx, r in enumerate(sorted_records, start=3):
        ws.row_dimensions[row_idx].height = 18

        total_orders_formula = f"=SUM(M{row_idx}:R{row_idx})"
        total_rev_formula = f"=SUM(U{row_idx}:Z{row_idx})"
        outstanding_a_formula = f"=AA{row_idx}-AB{row_idx}"
        total_excl_tax_formula = f"=AO{row_idx}+AQ{row_idx}+AS{row_idx}+AU{row_idx}+AW{row_idx}+AY{row_idx}+BA{row_idx}"
        total_sst_formula = f"=AP{row_idx}+AR{row_idx}+AT{row_idx}+AV{row_idx}+AX{row_idx}+AZ{row_idx}+BB{row_idx}"
        total_incl_tax_c_formula = f"=BC{row_idx}+BD{row_idx}"
        sales_incl_sst_formula = f"=BF{row_idx}+BG{row_idx}"
        todays_earnings_formula = f"=AC{row_idx}-BE{row_idx}"
        total_payable_formula = f"=BJ{row_idx}+BI{row_idx}"

        row_values = [
            r.get("data_source", "PDF only"),                       # 0 Col A
            r.get("business_date"),                                 # 1 Col B
            r.get("invoice_issue_date"),                            # 2 Col C
            r.get("period_from"),                                   # 3 Col D
            r.get("period_to"),                                     # 4 Col E
            r.get("outlet_name"),                                   # 5 Col F
            r.get("legal_entity"),                                  # 6 Col G
            r.get("invoice_no"),                                    # 7 Col H
            r.get("self_billed_inv_no"),                             # 8 Col I
            r.get("vendor_id"),                                     # 9 Col J
            r.get("vendor_code") or None,                           # 10 Col K
            r.get("vendor_tin") or None,                            # 11 Col L
            r.get("normal_orders", 0),                              # 12 Col M
            r.get("cancelled_partial_orders", 0),                   # 13 Col N
            r.get("cancelled_full_orders", 0),                      # 14 Col O
            r.get("pandabox_orders", 0),                            # 15 Col P
            r.get("dine_in_orders", 0),                             # 16 Col Q
            r.get("targeted_customer_orders", 0),                   # 17 Col R
            total_orders_formula,                                   # 18 Col S
            None,                                                   # 19 Col T
            r.get("normal_revenue", 0.0),                           # 20 Col U
            r.get("cancelled_partial_revenue", 0.0),                # 21 Col V
            r.get("cancelled_full_revenue", 0.0),                   # 22 Col W
            r.get("pandabox_revenue", 0.0),                         # 23 Col X
            r.get("dine_in_revenue", 0.0),                          # 24 Col Y
            r.get("targeted_customer_revenue", 0.0),                # 25 Col Z
            total_rev_formula,                                      # 26 Col AA
            r.get("already_received", 0.0),                         # 27 Col AB
            outstanding_a_formula,                                  # 28 Col AC
            None, None, None, None, None, None, None, None, None,   # 29-37 Col AD-AL
            None,                                                   # 38 Col AM
            r.get("commission_rate", 20.0),                         # 39 Col AN
            r.get("commission_base", 0.0),                          # 40 Col AO
            r.get("commission_sst", 0.0),                           # 41 Col AP
            r.get("wastage_commission_base", 0.0),                  # 42 Col AQ
            r.get("wastage_commission_sst", 0.0),                   # 43 Col AR
            r.get("fees_adj_sst_base", 0.0),                        # 44 Col AS
            r.get("fees_adj_sst_sst", 0.0),                         # 45 Col AT
            r.get("fees_adj_non_sst_base", 0.0),                    # 46 Col AU
            r.get("fees_adj_non_sst_sst", 0.0),                     # 47 Col AV
            r.get("pandabox_fee_base", 0.0),                        # 48 Col AW
            r.get("pandabox_fee_sst", 0.0),                         # 49 Col AX
            r.get("targeted_cust_fee_base", 0.0),                   # 50 Col AY
            r.get("targeted_cust_fee_sst", 0.0),                    # 51 Col AZ
            r.get("waiting_time_fee_base", 0.0),                    # 52 Col BA
            r.get("waiting_time_fee_sst", 0.0),                     # 53 Col BB
            total_excl_tax_formula,                                 # 54 Col BC
            total_sst_formula,                                      # 55 Col BD
            total_incl_tax_c_formula,                               # 56 Col BE
            r.get("sales_value_excl_sst", 0.0),                     # 57 Col BF
            r.get("sales_sst_6pct", 0.0),                           # 58 Col BG
            sales_incl_sst_formula,                                 # 59 Col BH
            todays_earnings_formula,                                # 60 Col BI
            r.get("outstanding_prev_f", 0.0),                       # 61 Col BJ
            total_payable_formula,                                  # 62 Col BK
            0,                                                      # 63 Col BL
            None,                                                   # 64 Col BM
            r.get("appendix_c_open_items", 0),                      # 65 Col BN
            r.get("source_file"),                                   # 66 Col BO
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row_idx, col_idx, val)
            cell.font = font_data
            cell.border = border_thin

            if isinstance(val, (int, float)) or (isinstance(val, str) and val.startswith("=")):
                if col_idx in [13, 14, 15, 16, 17, 18, 19, 64, 66]:
                    cell.number_format = "#,##0"
                    cell.alignment = align_right
                elif col_idx == 40:
                    cell.number_format = "0.00"
                    cell.alignment = align_right
                elif col_idx in list(range(21, 29)) + list(range(41, 64)):
                    cell.number_format = "#,##0.00"
                    cell.alignment = align_right
                else:
                    cell.alignment = align_right
            elif col_idx in [2, 3, 4, 5]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col[1:]:
            if cell.value:
                val_str = str(cell.value)
                if not val_str.startswith("="):
                    max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    wb.save(output_path)
    wb.close()


# ===========================================================================
# 9. Main Pipeline
# ===========================================================================

def main():
    parser = argparse.ArgumentParser(description="Transform Foodpanda PDF invoices to standardized Excel.")
    parser.add_argument("--input-dir", "-i", required=True, help="Directory containing Foodpanda PDF invoices.")
    parser.add_argument("--output", "-o", default="Foodpanda_Invoice_Report.xlsx", help="Output .xlsx file path.")
    parser.add_argument("--workers", "-w", type=int, default=os.cpu_count() or 4, help="Number of worker processes.")
    parser.add_argument("--limit", "-l", type=int, default=None, help="Limit number of PDFs to process (for testing).")
    parser.add_argument("--skip-sister-brands", action="store_true", help="Exclude non-US Pizza brands.")

    args = parser.parse_args()

    if not os.path.isdir(args.input_dir):
        sys.exit(f"Error: Directory '{args.input_dir}' does not exist.")

    pdf_pattern = os.path.join(args.input_dir, "*.pdf")
    pdf_files = sorted(glob.glob(pdf_pattern))
    if not pdf_files:
        pdf_pattern = os.path.join(args.input_dir, "**", "*.pdf")
        pdf_files = sorted(glob.glob(pdf_pattern, recursive=True))

    if not pdf_files:
        sys.exit(f"No PDF files found in '{args.input_dir}'.")

    if args.limit:
        pdf_files = pdf_files[:args.limit]

    print(f"============================================================")
    print(f" Foodpanda Invoice PDF -> Excel Transformer")
    print(f" Input:   {args.input_dir} ({len(pdf_files)} PDF files)")
    print(f" Output:  {args.output}")
    print(f" Workers: {args.workers}")
    print(f"============================================================")

    start_time = datetime.datetime.now()

    records = []
    errors = []
    with concurrent.futures.ProcessPoolExecutor(max_workers=args.workers) as executor:
        futures = {executor.submit(parse_single_pdf, path): path for path in pdf_files}
        for future in concurrent.futures.as_completed(futures):
            res = future.result()
            if res.get("error"):
                errors.append((res.get("source_file"), res.get("error")))
            else:
                if args.skip_sister_brands:
                    name = res.get("outlet_name", "").lower()
                    if "manhattan" in name or "ny steak" in name:
                        continue
                records.append(res)

    duration = (datetime.datetime.now() - start_time).total_seconds()
    print(f"\n[DONE] Parsed {len(records)} invoices in {duration:.2f} seconds ({len(records)/max(duration, 0.01):.1f} files/sec).")
    if errors:
        print(f"[WARN] {len(errors)} files failed to parse:")
        for fn, err in errors[:5]:
            print(f"   - {fn}: {err}")

    # Write Excel
    print(f"[INFO] Generating Excel workbook: {args.output}...")
    write_excel_report(records, args.output)
    print(f"[SUCCESS] Wrote '{args.output}' with 'Invoice Master' sheet.")

    # Summary
    tot_orders = sum(r.get("total_orders", 0) for r in records)
    tot_rev = sum(r.get("total_revenue", 0.0) for r in records)
    tot_comm_base = sum(r.get("commission_base", 0.0) for r in records)
    tot_comm_sst = sum(r.get("commission_sst", 0.0) for r in records)
    tot_wastage_base = sum(r.get("wastage_commission_base", 0.0) for r in records)
    tot_wastage_sst = sum(r.get("wastage_commission_sst", 0.0) for r in records)
    tot_sales_excl_sst = sum(r.get("sales_value_excl_sst", 0.0) for r in records)
    tot_sales_sst_6 = sum(r.get("sales_sst_6pct", 0.0) for r in records)
    tot_earnings = sum(r.get("todays_earnings", 0.0) for r in records)
    tot_payable = sum(r.get("total_payable", 0.0) for r in records)

    print(f"\n------------------------------------------------------------")
    print(f" BATCH RECONCILIATION SUMMARY ({len(records)} invoices)")
    print(f"------------------------------------------------------------")
    print(f" Total Completed Orders:       {tot_orders:,}")
    print(f" Total Revenue (Sales Incl):   RM {tot_rev:,.2f}")
    print(f" Sales Value Excl. SST (Net):  RM {tot_sales_excl_sst:,.2f}")
    print(f" Customer Sales Tax (6% SST):  RM {tot_sales_sst_6:,.2f}")
    print(f" Commission Base:              RM {tot_comm_base:,.2f}")
    print(f" Commission 8% SST:            RM {tot_comm_sst:,.2f}")
    print(f" Wastage Commission (Base+SST):RM {tot_wastage_base + tot_wastage_sst:,.2f}")
    print(f" Today's Net Earnings:         RM {tot_earnings:,.2f}")
    print(f" Total Bank Payable:           RM {tot_payable:,.2f}")
    print(f"------------------------------------------------------------\n")


if __name__ == "__main__":
    main()
